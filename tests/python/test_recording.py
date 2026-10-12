import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from serial_ledger.cli import demo
from serial_ledger.export import export_xlsx
from serial_ledger.recording import MAX_LINE_BYTES, LineFramer, Recorder, parse_reading


class RecordingTests(unittest.TestCase):
    def test_split_timeout_and_multiple_lines(self):
        framer = LineFramer()
        self.assertEqual(list(framer.feed(b'{"a":')), [])
        self.assertEqual(list(framer.feed(b'1}\r\n{"a":2}\n')), [b'{"a":1}', b'{"a":2}'])

    def test_oversized_line_recovers(self):
        framer = LineFramer()
        self.assertEqual(list(framer.feed(b"x" * (MAX_LINE_BYTES + 2))), [])
        self.assertEqual(list(framer.feed(b'\n{"a":1}\n')), [b'{"a":1}'])
        self.assertEqual(framer.dropped, 1)

    def test_nonfinite_and_nested_values_remain_raw(self):
        for text in ['{"a":NaN}', '{"a":1e999}', '{"a":{}}', "[]", "debug text"]:
            self.assertIsNone(parse_reading(text))

    def test_excel_types_raw_messages_and_formula_text(self):
        with tempfile.TemporaryDirectory() as directory:
            csv_path = Path(directory) / "session.csv"
            xlsx_path = Path(directory) / "session.xlsx"
            with Recorder(csv_path) as recorder:
                recorder.append(b'{"temperature_C":25.4,"label":"=1+1","enabled":true}')
                recorder.append(b"debug text")
                recorder.append(b'{"temperature_C":26,"humidity_pct":60}')
            self.assertEqual(export_xlsx(csv_path, xlsx_path), 3)
            workbook = load_workbook(xlsx_path)
            try:
                sheet = workbook["Readings"]
                rows = list(sheet.iter_rows(values_only=True))
                positions = {key: index for index, key in enumerate(rows[0])}
                self.assertEqual(rows[1][positions["value.temperature_C"]], 25.4)
                self.assertEqual(rows[1][positions["value.enabled"]], True)
                self.assertEqual(rows[2][positions["raw"]], "debug text")
                self.assertEqual(rows[2][positions["parse_status"]], "raw")
                label = sheet.cell(2, positions["value.label"] + 1)
                self.assertEqual(label.value, "=1+1")
                self.assertEqual(label.data_type, "s")
            finally:
                workbook.close()

    def test_session_file_is_not_overwritten(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "session.csv"
            path.write_text("existing", encoding="utf-8")
            with self.assertRaises(FileExistsError):
                Recorder(path)
            self.assertEqual(path.read_text(encoding="utf-8"), "existing")

    def test_mixed_serial_output_preserves_raw_and_recovers_in_export(self):
        lines = [
            b'{"seq":0,"time_ms":1000}',
            b"[DEBUG] mixed-output test",
            b'{"temperature_C":25.4',
            b'{"temperature_C":26.5,"humidity_pct":60}',
        ]
        wire = b"\r\n".join(lines) + b"\r\n"
        with tempfile.TemporaryDirectory() as directory:
            csv_path = Path(directory) / "mixed.csv"
            xlsx_path = Path(directory) / "mixed.xlsx"
            framer = LineFramer()
            with Recorder(csv_path) as recorder:
                for chunk in [wire[:5], wire[5:59], wire[59:]]:
                    for line in framer.feed(chunk):
                        recorder.append(line)
            with csv_path.open(newline="", encoding="utf-8") as stream:
                captured = list(csv.DictReader(stream))
            self.assertEqual([row["raw"] for row in captured], [line.decode() for line in lines])
            self.assertEqual([row["parse_status"] for row in captured], ["ok", "raw", "raw", "ok"])
            self.assertEqual(export_xlsx(csv_path, xlsx_path), 4)
            workbook = load_workbook(xlsx_path)
            try:
                rows = list(workbook["Readings"].iter_rows(values_only=True))
                positions = {key: index for index, key in enumerate(rows[0])}
                self.assertEqual(rows[1][positions["value.seq"]], 0)
                for row, expected in zip(rows[2:4], lines[1:3]):
                    self.assertEqual(row[positions["raw"]], expected.decode())
                    self.assertEqual(row[positions["parse_status"]], "raw")
                    self.assertTrue(all(value is None for value in row[3:]))
                self.assertEqual(rows[4][positions["value.temperature_C"]], 26.5)
                self.assertEqual(rows[4][positions["value.humidity_pct"]], 60)
            finally:
                workbook.close()

    def test_loopback_demo_exports_three_samples(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "demo.csv"
            demo(path)
            with path.open(newline="", encoding="utf-8") as stream:
                self.assertEqual(len(list(csv.DictReader(stream))), 3)
            self.assertTrue(path.with_suffix(".xlsx").exists())


if __name__ == "__main__":
    unittest.main()
